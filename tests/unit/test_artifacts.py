import io
import docx
import openpyxl
import pptx
import zipfile

from packages.artifact_tools.generator import ArtifactGenerator

def test_generate_pdf():
    pdf_bytes = ArtifactGenerator.generate_pdf("Test Title", "Line 1\nLine 2")
    assert pdf_bytes is not None
    assert len(pdf_bytes) > 100
    assert pdf_bytes.startswith(b"%PDF-")

def test_generate_docx():
    docx_bytes = ArtifactGenerator.generate_docx("Sample Document", "Paragraph 1\nParagraph 2")
    assert docx_bytes is not None
    assert len(docx_bytes) > 500
    # Reopen and parse with python-docx to verify document structure integrity
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs]
    assert "Sample Document" in paragraphs[0]
    assert "Paragraph 1" in paragraphs

def test_generate_xlsx():
    headers = ["Task", "Cost", "Completed"]
    rows = [["Design", 150, True], ["Build", 300, False]]
    xlsx_bytes = ArtifactGenerator.generate_xlsx("BudgetSheet", headers, rows)
    assert xlsx_bytes is not None
    # Reopen with openpyxl to verify formula/table readability
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = wb["BudgetSheet"]
    assert sheet.cell(row=1, column=1).value == "Task"
    assert sheet.cell(row=2, column=2).value == 150

def test_generate_pptx():
    bullets = ["Point A", "Point B", "Point C"]
    pptx_bytes = ArtifactGenerator.generate_pptx("Presentation 1", "Subtitle", bullets)
    assert pptx_bytes is not None
    # Reopen with python-pptx
    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title.text == "Presentation 1"

def test_generate_zip():
    files = {
        "hello.txt": b"Hello World",
        "data.json": b'{"key": "value"}'
    }
    zip_bytes = ArtifactGenerator.generate_zip(files)
    assert zip_bytes is not None
    # Verify zip extraction integrity
    with zipfile.ZipFile(io.BytesIO(zip_bytes), "r") as zf:
        namelist = zf.namelist()
        assert "hello.txt" in namelist
        assert zf.read("hello.txt") == b"Hello World"
